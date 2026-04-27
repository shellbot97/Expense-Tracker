"""
Unit tests for Excel parser
"""

import pytest
from io import BytesIO
import openpyxl
from datetime import datetime
from src.parsers.excel_parser import ExcelParser, EXCEL_SUPPORT
from src.utils.exceptions import ValidationError


def create_test_excel(data, headers=None):
    """Helper function to create test Excel file"""
    wb = openpyxl.Workbook()
    ws = wb.active
    
    if headers:
        ws.append(headers)
    
    for row in data:
        ws.append(row)
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


# ============= Basic Parsing Tests =============

@pytest.mark.unit
@pytest.mark.skipif(not EXCEL_SUPPORT, reason="openpyxl not installed")
class TestExcelParserBasic:
    """Tests for basic Excel parsing"""

    def test_parse_simple_excel(self):
        """Test parsing simple Excel with default mapping"""
        data = [
            ("2024-01-15", "Walmart", 50.00, "expense"),
            ("2024-01-16", "Paycheck", 1500.00, "income"),
        ]
        headers = ["date", "description", "amount", "transaction_type"]
        
        excel_bytes = create_test_excel(data, headers)
        
        parser = ExcelParser(column_mapping={
            "date": "date",
            "description": "description",
            "amount": "amount",
            "transaction_type": "transaction_type",
        })
        
        transactions = parser.parse(excel_bytes)
        
        assert len(transactions) == 2
        assert transactions[0]["date"] == "2024-01-15"
        assert transactions[0]["description"] == "Walmart"
        assert transactions[0]["amount"] == 5000  # cents
        assert transactions[0]["transaction_type"] == "expense"

    def test_parse_empty_excel_raises_error(self):
        """Test that empty Excel raises ValidationError"""
        parser = ExcelParser()
        
        with pytest.raises(ValidationError, match="empty"):
            parser.parse(b"")

    def test_parse_no_transactions_raises_error(self):
        """Test that Excel with only header raises error"""
        headers = ["date", "description", "amount"]
        excel_bytes = create_test_excel([], headers)
        
        parser = ExcelParser(column_mapping={
            "date": "date",
            "description": "description",
            "amount": "amount",
        })
        
        with pytest.raises(ValidationError, match="No transactions found"):
            parser.parse(excel_bytes)


# ============= Date Parsing Tests =============

@pytest.mark.unit
@pytest.mark.skipif(not EXCEL_SUPPORT, reason="openpyxl not installed")
class TestExcelDateParsing:
    """Tests for date parsing"""

    def test_parse_datetime_object(self):
        """Test parsing datetime object"""
        parser = ExcelParser()
        dt = datetime(2024, 1, 15)
        date_str = parser._parse_date(dt)
        assert date_str == "2024-01-15"

    def test_parse_date_string(self):
        """Test parsing date string"""
        parser = ExcelParser()
        date_str = parser._parse_date("2024-01-15")
        assert date_str == "2024-01-15"

    def test_parse_us_date_string(self):
        """Test parsing US format date string"""
        parser = ExcelParser()
        date_str = parser._parse_date("01/15/2024")
        assert date_str == "2024-01-15"

    def test_parse_invalid_date_raises_error(self):
        """Test that invalid date raises error"""
        parser = ExcelParser()
        with pytest.raises(ValidationError, match="Could not parse date"):
            parser._parse_date("invalid-date")


# ============= Amount Parsing Tests =============

@pytest.mark.unit
@pytest.mark.skipif(not EXCEL_SUPPORT, reason="openpyxl not installed")
class TestExcelAmountParsing:
    """Tests for amount parsing"""

    def test_parse_float_amount(self):
        """Test parsing float amount"""
        parser = ExcelParser()
        amount = parser._parse_amount(100.50)
        assert amount == 10050

    def test_parse_integer_amount(self):
        """Test parsing integer amount"""
        parser = ExcelParser()
        amount = parser._parse_amount(100)
        assert amount == 10000

    def test_parse_string_amount(self):
        """Test parsing string amount"""
        parser = ExcelParser()
        amount = parser._parse_amount("100.50")
        assert amount == 10050

    def test_parse_amount_with_dollar_sign(self):
        """Test parsing amount with currency symbol"""
        parser = ExcelParser()
        amount = parser._parse_amount("$100.50")
        assert amount == 10050

    def test_parse_negative_amount(self):
        """Test parsing negative amount"""
        parser = ExcelParser()
        amount = parser._parse_amount(-100.50)
        assert amount == -10050

    def test_parse_invalid_amount_raises_error(self):
        """Test that invalid amount string raises error"""
        parser = ExcelParser()
        with pytest.raises(ValidationError, match="Could not parse amount"):
            parser._parse_amount("not-a-number")


# ============= Multi-Sheet Tests =============

@pytest.mark.unit
@pytest.mark.skipif(not EXCEL_SUPPORT, reason="openpyxl not installed")
class TestExcelMultiSheet:
    """Tests for multi-sheet Excel files"""

    def test_get_sheet_names(self):
        """Test getting sheet names"""
        wb = openpyxl.Workbook()
        wb.create_sheet("Sheet2")
        wb.create_sheet("Transactions")
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        excel_bytes = output.getvalue()
        
        parser = ExcelParser()
        sheet_names = parser.get_sheet_names(excel_bytes)
        
        assert "Sheet" in sheet_names
        assert "Sheet2" in sheet_names
        assert "Transactions" in sheet_names

    def test_parse_specific_sheet(self):
        """Test parsing specific sheet by name"""
        # Create Excel with multiple sheets
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Summary"
        ws1.append(["date", "description", "amount"])
        ws1.append(("2024-01-15", "Wrong sheet", 100.00))
        
        ws2 = wb.create_sheet("Transactions")
        ws2.append(["date", "description", "amount"])
        ws2.append(("2024-01-16", "Correct sheet", 200.00))
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        excel_bytes = output.getvalue()
        
        parser = ExcelParser(column_mapping={
            "date": "date",
            "description": "description",
            "amount": "amount",
        })
        
        transactions = parser.parse(excel_bytes, sheet_name="Transactions")
        
        assert len(transactions) == 1
        assert transactions[0]["description"] == "Correct sheet"

    def test_parse_invalid_sheet_name_raises_error(self):
        """Test that invalid sheet name raises error"""
        excel_bytes = create_test_excel(
            [("2024-01-15", "Test", 100.00)],
            ["date", "description", "amount"]
        )
        
        parser = ExcelParser()
        
        with pytest.raises(ValidationError, match="Sheet .* not found"):
            parser.parse(excel_bytes, sheet_name="NonExistent")


# ============= Column Detection Tests =============

@pytest.mark.unit
@pytest.mark.skipif(not EXCEL_SUPPORT, reason="openpyxl not installed")
class TestExcelColumnDetection:
    """Tests for column detection"""

    def test_detect_columns(self):
        """Test detecting Excel columns"""
        headers = ["Date", "Description", "Amount", "Type"]
        excel_bytes = create_test_excel(
            [("2024-01-15", "Test", 100.00, "expense")],
            headers
        )
        
        parser = ExcelParser()
        columns = parser.detect_columns(excel_bytes)
        
        assert columns == headers

    def test_detect_columns_empty_excel(self):
        """Test detecting columns from invalid Excel"""
        parser = ExcelParser()
        columns = parser.detect_columns(b"invalid-excel-data")
        assert columns == []


# ============= Error Handling Tests =============

@pytest.mark.unit
@pytest.mark.skipif(not EXCEL_SUPPORT, reason="openpyxl not installed")
class TestExcelErrorHandling:
    """Tests for error handling"""

    def test_parse_missing_required_field(self):
        """Test parsing Excel with missing required field"""
        excel_bytes = create_test_excel(
            [("Walmart", "expense")],
            ["description", "transaction_type"]
        )
        
        parser = ExcelParser(column_mapping={
            "description": "description",
            "transaction_type": "transaction_type",
        })
        
        with pytest.raises(ValidationError, match="Missing required field"):
            parser.parse(excel_bytes)

    def test_parse_invalid_excel_raises_error(self):
        """Test parsing invalid Excel file"""
        parser = ExcelParser()
        
        with pytest.raises(ValidationError, match="Error parsing Excel"):
            parser.parse(b"not-an-excel-file")

    def test_parse_skips_empty_rows(self):
        """Test that empty rows are skipped"""
        data = [
            ("2024-01-15", "First", 100.00),
            (None, None, None),  # Empty row
            ("2024-01-16", "Second", 200.00),
        ]
        excel_bytes = create_test_excel(data, ["date", "description", "amount"])
        
        parser = ExcelParser(column_mapping={
            "date": "date",
            "description": "description",
            "amount": "amount",
        })
        
        transactions = parser.parse(excel_bytes)
        
        # Should only have 2 transactions (empty row skipped)
        assert len(transactions) == 2
        assert transactions[0]["description"] == "First"
        assert transactions[1]["description"] == "Second"
