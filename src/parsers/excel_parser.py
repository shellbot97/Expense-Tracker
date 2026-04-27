"""
Excel statement parser for importing bank transactions
"""

from typing import List, Dict, Optional
from io import BytesIO
from datetime import datetime
from src.utils.exceptions import ValidationError

try:
    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException
    EXCEL_SUPPORT = True
except ImportError:
    EXCEL_SUPPORT = False


class ExcelParser:
    """
    Parse Excel bank statements into structured transaction data
    
    Supports flexible column mapping to handle different bank formats.
    Requires openpyxl library.
    """
    
    def __init__(self, column_mapping: Optional[Dict[str, str]] = None):
        """
        Initialize Excel parser with column mapping
        
        Args:
            column_mapping: Dict mapping Excel columns to transaction fields
                Example: {"Date": "date", "Description": "description", "Amount": "amount"}
                If None, uses default mapping
        """
        if not EXCEL_SUPPORT:
            raise ValidationError("Excel support not available. Install openpyxl: pip install openpyxl")
        
        self.column_mapping = column_mapping or {
            "date": "date",
            "description": "description",
            "amount": "amount",
            "transaction_type": "transaction_type",
        }
    
    def parse(self, file_bytes: bytes, sheet_name: Optional[str] = None, has_header: bool = True) -> List[Dict]:
        """
        Parse Excel content into transaction dictionaries
        
        Args:
            file_bytes: Raw Excel file bytes
            sheet_name: Sheet name to parse (defaults to first sheet)
            has_header: Whether Excel has header row
            
        Returns:
            List of transaction dictionaries
            
        Raises:
            ValidationError: If Excel format is invalid
        """
        if not file_bytes:
            raise ValidationError("Excel content is empty")
        
        if not EXCEL_SUPPORT:
            raise ValidationError("Excel support not available. Install openpyxl: pip install openpyxl")
        
        try:
            # Load workbook
            workbook = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
            
            # Select sheet
            if sheet_name:
                if sheet_name not in workbook.sheetnames:
                    raise ValidationError(f"Sheet '{sheet_name}' not found")
                sheet = workbook[sheet_name]
            else:
                sheet = workbook.active
            
            # Parse rows
            transactions = []
            rows = list(sheet.iter_rows(values_only=True))
            
            if not rows:
                raise ValidationError("Excel sheet is empty")
            
            # Extract header if present
            if has_header:
                header = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
                data_rows = rows[1:]
            else:
                # Generate default header
                header = [f"Column{i+1}" for i in range(len(rows[0]))]
                data_rows = rows
            
            # Parse each row
            for row_num, row in enumerate(data_rows, start=2 if has_header else 1):
                if all(cell is None or str(cell).strip() == "" for cell in row):
                    # Skip empty rows
                    continue
                
                try:
                    row_dict = dict(zip(header, row))
                    transaction = self._parse_row(row_dict)
                    transactions.append(transaction)
                except Exception as e:
                    raise ValidationError(f"Error parsing row {row_num}: {str(e)}")
            
            if not transactions:
                raise ValidationError("No transactions found in Excel")
            
            workbook.close()
            return transactions
            
        except InvalidFileException:
            raise ValidationError("Invalid Excel file format")
        except Exception as e:
            if isinstance(e, ValidationError):
                raise
            raise ValidationError(f"Error parsing Excel: {str(e)}")
    
    def _parse_row(self, row: Dict[str, any]) -> Dict:
        """
        Parse a single Excel row into transaction dict
        
        Args:
            row: Excel row as dictionary
            
        Returns:
            Transaction dictionary
        """
        transaction = {}
        
        # Map columns
        for excel_col, field_name in self.column_mapping.items():
            if excel_col in row:
                value = row[excel_col]
                
                # Skip None values
                if value is None:
                    continue
                
                # Parse based on field type
                if field_name == "date":
                    transaction[field_name] = self._parse_date(value)
                elif field_name == "amount":
                    transaction[field_name] = self._parse_amount(value)
                else:
                    transaction[field_name] = str(value).strip()
        
        # Validate required fields
        required_fields = ["date", "amount"]
        for field in required_fields:
            if field not in transaction:
                raise ValidationError(f"Missing required field: {field}")
        
        return transaction
    
    def _parse_date(self, date_value) -> str:
        """
        Parse date value into ISO format (YYYY-MM-DD)
        
        Handles datetime objects and string dates
        """
        # If already a datetime object
        if isinstance(date_value, datetime):
            return date_value.strftime("%Y-%m-%d")
        
        # If string, try to parse
        if isinstance(date_value, str):
            date_str = date_value.strip()
            
            if not date_str:
                raise ValidationError("Date is empty")
            
            # Common date formats
            formats = [
                "%Y-%m-%d",      # 2024-01-15
                "%m/%d/%Y",      # 01/15/2024
                "%d/%m/%Y",      # 15/01/2024
                "%Y/%m/%d",      # 2024/01/15
                "%m-%d-%Y",      # 01-15-2024
                "%d-%m-%Y",      # 15-01-2024
                "%b %d, %Y",     # Jan 15, 2024
                "%B %d, %Y",     # January 15, 2024
                "%d %b %Y",      # 15 Jan 2024
                "%d %B %Y",      # 15 January 2024
            ]
            
            for fmt in formats:
                try:
                    dt = datetime.strptime(date_str, fmt)
                    return dt.strftime("%Y-%m-%d")
                except ValueError:
                    continue
            
            raise ValidationError(f"Could not parse date: {date_str}")
        
        raise ValidationError(f"Invalid date type: {type(date_value)}")
    
    def _parse_amount(self, amount_value) -> int:
        """
        Parse amount value into cents (integer)
        
        Handles numbers and strings
        """
        # If already a number
        if isinstance(amount_value, (int, float)):
            amount_cents = int(round(amount_value * 100))
            return amount_cents
        
        # If string, parse like CSV
        if isinstance(amount_value, str):
            amount_str = amount_value.strip()
            
            if not amount_str:
                raise ValidationError("Amount is empty")
            
            # Remove currency symbols and spaces
            amount_str = amount_str.replace("$", "").replace("€", "").replace("£", "")
            amount_str = amount_str.replace(",", "").replace(" ", "")
            
            # Handle parentheses as negative
            is_negative = False
            if amount_str.startswith("(") and amount_str.endswith(")"):
                is_negative = True
                amount_str = amount_str[1:-1]
            
            try:
                amount_float = float(amount_str)
                
                if is_negative:
                    amount_float = -amount_float
                
                amount_cents = int(round(amount_float * 100))
                return amount_cents
                
            except ValueError:
                raise ValidationError(f"Could not parse amount: {amount_str}")
        
        raise ValidationError(f"Invalid amount type: {type(amount_value)}")
    
    def get_sheet_names(self, file_bytes: bytes) -> List[str]:
        """
        Get list of sheet names from Excel file
        
        Args:
            file_bytes: Raw Excel file bytes
            
        Returns:
            List of sheet names
        """
        if not EXCEL_SUPPORT:
            raise ValidationError("Excel support not available")
        
        try:
            workbook = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
            sheet_names = workbook.sheetnames
            workbook.close()
            return sheet_names
        except Exception:
            return []
    
    def detect_columns(self, file_bytes: bytes, sheet_name: Optional[str] = None) -> List[str]:
        """
        Detect column headers from Excel file
        
        Args:
            file_bytes: Raw Excel file bytes
            sheet_name: Sheet name (defaults to first sheet)
            
        Returns:
            List of column names
        """
        if not EXCEL_SUPPORT:
            raise ValidationError("Excel support not available")
        
        try:
            workbook = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
            
            if sheet_name:
                sheet = workbook[sheet_name]
            else:
                sheet = workbook.active
            
            # Get first row
            first_row = next(sheet.iter_rows(values_only=True), None)
            workbook.close()
            
            if first_row:
                return [str(cell).strip() if cell is not None else "" for cell in first_row]
            return []
            
        except Exception:
            return []
